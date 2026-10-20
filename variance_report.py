"""Budget vs Actual variance report.

Reads a monthly budget and posted G/L entries (Business Central layout),
computes variances by account and cost center and writes a formatted
Excel report.

Usage:
    python variance_report.py                    # latest month in the data
    python variance_report.py --month 2026-06    # month-to-date + YTD up to June
    python variance_report.py --threshold 0.10   # flag variances above 10%
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
DATA_DIR = BASE / "data"
OUTPUT_DIR = BASE / "output"

KEYS = ["G/L Account No.", "G/L Account Name", "Account Type"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BFBFBF")


# ---------------------------------------------------------------- data ----
def load_data() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not (DATA_DIR / "budget.csv").exists():
        import generate_sample_data

        generate_sample_data.main()

    budget = pd.read_csv(DATA_DIR / "budget.csv", dtype={"G/L Account No.": str})
    actuals = pd.read_csv(DATA_DIR / "actuals.csv", dtype={"G/L Account No.": str})

    budget["Period"] = pd.to_datetime(budget["Period"]).dt.to_period("M")
    actuals["Period"] = pd.to_datetime(actuals["Posting Date"]).dt.to_period("M")

    # Attach account master data to G/L entries and flip revenue sign
    # (BC posts revenue as negative amounts).
    accounts = budget[KEYS].drop_duplicates()
    actuals = actuals.merge(accounts, on="G/L Account No.", how="left")
    actuals["Actual Amount"] = actuals["Amount"].where(
        actuals["Account Type"] != "Revenue", -actuals["Amount"]
    )
    return budget, actuals


def variance_table(budget: pd.DataFrame, actuals: pd.DataFrame, by: list[str]) -> pd.DataFrame:
    b = budget.groupby(by, as_index=False)["Budget Amount"].sum()
    a = actuals.groupby(by, as_index=False)["Actual Amount"].sum()
    df = b.merge(a, on=by, how="outer").fillna({"Budget Amount": 0, "Actual Amount": 0})
    df["Variance"] = df["Actual Amount"] - df["Budget Amount"]
    df["Variance %"] = df["Variance"] / df["Budget Amount"].where(df["Budget Amount"] != 0)
    if "Account Type" in df.columns:
        # Favourable: more revenue or less cost than budget
        sign = df["Account Type"].map({"Revenue": 1, "Cost": -1})
        df["Favourable"] = (df["Variance"] * sign) >= 0
    return df.sort_values(by).reset_index(drop=True)


def operating_result(df: pd.DataFrame) -> dict:
    rev = df[df["Account Type"] == "Revenue"]
    cost = df[df["Account Type"] == "Cost"]
    out = {}
    for col in ["Budget Amount", "Actual Amount"]:
        out[col] = rev[col].sum() - cost[col].sum()
    out["Variance"] = out["Actual Amount"] - out["Budget Amount"]
    out["Variance %"] = out["Variance"] / out["Budget Amount"] if out["Budget Amount"] else None
    return out


# --------------------------------------------------------------- excel ----
def write_table(ws, df: pd.DataFrame, start_row: int, threshold: float, title: str) -> int:
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=12)
    start_row += 1
    cols = [c for c in df.columns if c != "Favourable"]
    for j, col in enumerate(cols, 1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        rec = dict(zip(df.columns, row))
        for j, col in enumerate(cols, 1):
            val = rec[col]
            if isinstance(val, pd.Period):
                val = str(val)
            c = ws.cell(row=i, column=j, value=None if pd.isna(val) else val)
            c.border = Border(bottom=THIN)
            if col.endswith("Amount") or col == "Variance":
                c.number_format = "#,##0;[Red]-#,##0"
            elif col == "Variance %":
                c.number_format = "0.0%;[Red]-0.0%"
        # Highlight rows whose variance exceeds the threshold
        if "Favourable" in rec and pd.notna(rec["Variance %"]) and abs(rec["Variance %"]) >= threshold:
            fill = GREEN_FILL if rec["Favourable"] else RED_FILL
            for j in range(1, len(cols) + 1):
                ws.cell(row=i, column=j).fill = fill
    return start_row + len(df) + 2


def autosize(ws) -> None:
    for col_cells in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width + 2, 10), 40)


def build_report(month: pd.Period, threshold: float) -> Path:
    budget, actuals = load_data()
    fy_start = budget["Period"].min()

    b_mtd, a_mtd = budget[budget["Period"] == month], actuals[actuals["Period"] == month]
    b_ytd = budget[(budget["Period"] >= fy_start) & (budget["Period"] <= month)]
    a_ytd = actuals[(actuals["Period"] >= fy_start) & (actuals["Period"] <= month)]

    mtd = variance_table(b_mtd, a_mtd, KEYS)
    ytd = variance_table(b_ytd, a_ytd, KEYS)
    by_cc = variance_table(b_ytd, a_ytd, ["Cost Center Code"] + KEYS)
    trend = variance_table(budget[budget["Period"] <= month], actuals[actuals["Period"] <= month],
                           ["Period", "Account Type"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Budget vs Actual - {month.strftime('%B %Y')}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Fiscal year from {fy_start.strftime('%b %Y')} | rows highlighted when |variance| >= {threshold:.0%}"
    ws["A2"].font = Font(italic=True, color="595959")

    # KPI block: operating result (revenue - costs)
    for j, key in enumerate(["Budget", "Actual", "Variance", "Variance %"], 2):
        c = ws.cell(row=3, column=j, value=key)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    r = 4
    for label, data in [("Month", mtd), ("Year to date", ytd)]:
        res = operating_result(data)
        ws.cell(row=r, column=1, value=f"Operating result - {label}").font = TOTAL_FONT
        for j, key in enumerate(["Budget Amount", "Actual Amount", "Variance", "Variance %"], 2):
            c = ws.cell(row=r, column=j, value=res[key])
            c.number_format = "0.0%;[Red]-0.0%" if key == "Variance %" else "#,##0;[Red]-#,##0"
            c.font = TOTAL_FONT
        r += 1

    r += 1
    r = write_table(ws, mtd, r, threshold, "Month")
    write_table(ws, ytd, r, threshold, "Year to date")
    autosize(ws)

    ws2 = wb.create_sheet("By cost center")
    write_table(ws2, by_cc, 1, threshold, "Year to date by cost center")
    autosize(ws2)

    ws3 = wb.create_sheet("Monthly trend")
    write_table(ws3, trend, 1, threshold, "Monthly trend by account type")
    autosize(ws3)

    # Top deviations (absolute, YTD)
    ws4 = wb.create_sheet("Top deviations")
    top = by_cc.reindex(by_cc["Variance"].abs().sort_values(ascending=False).index).head(10)
    write_table(ws4, top, 1, threshold, "Top 10 YTD deviations (account x cost center)")
    autosize(ws4)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False

    OUTPUT_DIR.mkdir(exist_ok=True)
    out = OUTPUT_DIR / f"variance_report_{month}.xlsx"
    wb.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--month", help="Reporting month, YYYY-MM (default: latest month with actuals)")
    parser.add_argument("--threshold", type=float, default=0.10, help="Variance %% to highlight (default 0.10)")
    args = parser.parse_args()

    if args.month:
        month = pd.Period(args.month, freq="M")
    else:
        _, actuals = load_data()
        month = actuals["Period"].max()

    out = build_report(month, args.threshold)
    print(f"Report written to {out}")


if __name__ == "__main__":
    main()
